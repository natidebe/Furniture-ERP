from decimal import Decimal

import pytest
from django.core.management import CommandError, call_command
from openpyxl import Workbook

from apps.catalog.models import PriceHistory, Product


def _xlsx(tmp_path, rows):
    wb = Workbook()
    ws = wb.active
    ws.append(["Code", "Name", "Category", "Unit", "Price"])
    for row in rows:
        ws.append(row)
    path = tmp_path / "products.xlsx"
    wb.save(path)
    return str(path)


@pytest.mark.django_db
def test_import_creates_products(tmp_path):
    path = _xlsx(tmp_path, [
        ["vc-001", "Visitor chair", "Visitor chairs", "pcs", 2500],
        ["DS-003", "Manager desk", "Desks", "Pieces", "18000.50"],
    ])

    call_command("import_products", path)

    assert Product.objects.get(code="VC-001").selling_price == Decimal("2500.00")
    assert Product.objects.get(code="DS-003").unit.symbol == "pcs"


@pytest.mark.django_db
def test_import_is_idempotent_and_records_price_changes(tmp_path):
    call_command("import_products",
                 _xlsx(tmp_path, [["VC-001", "Visitor chair", "Visitor chairs", "pcs", 2500]]))
    call_command("import_products",
                 _xlsx(tmp_path, [["VC-001", "Visitor chair", "Visitor chairs", "pcs", 2500]]))
    assert Product.objects.count() == 1
    assert not PriceHistory.objects.exists()

    call_command("import_products",
                 _xlsx(tmp_path, [["VC-001", "Visitor chair", "Visitor chairs", "pcs", 2600]]))

    assert Product.objects.get(code="VC-001").selling_price == Decimal("2600.00")
    assert PriceHistory.objects.get().reason == "import_products"


@pytest.mark.django_db
def test_dry_run_saves_nothing(tmp_path):
    call_command("import_products",
                 _xlsx(tmp_path, [["VC-001", "Visitor chair", "Visitor chairs", "pcs", 2500]]),
                 dry_run=True)

    assert not Product.objects.exists()


@pytest.mark.django_db
def test_any_bad_row_aborts_whole_import(tmp_path):
    path = _xlsx(tmp_path, [
        ["VC-001", "Visitor chair", "Visitor chairs", "pcs", 2500],
        ["VC-002", "Bad unit", "Visitor chairs", "boxes", 100],
        ["VC-003", "Bad price", "Visitor chairs", "pcs", 0],
    ])

    with pytest.raises(CommandError) as exc:
        call_command("import_products", path)

    assert "row 3" in str(exc.value) and "row 4" in str(exc.value)
    assert not Product.objects.exists()


@pytest.mark.django_db
def test_optional_min_stock_and_description(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.append(["Code", "Name", "Category", "Unit", "Price", "Min stock", "Description"])
    ws.append(["VC-001", "Visitor chair", "Visitor chairs", "pcs", 2500, 10, "Black mesh"])
    path = tmp_path / "products.xlsx"
    wb.save(path)

    call_command("import_products", str(path))

    product = Product.objects.get(code="VC-001")
    assert (product.min_stock, product.description) == (10, "Black mesh")


@pytest.mark.django_db
def test_category_matched_case_insensitively(tmp_path):
    from apps.catalog.models import Category

    before = Category.objects.count()
    call_command("import_products",
                 _xlsx(tmp_path, [["VC-001", "Visitor chair", "visitor CHAIRS", "pcs", 2500]]))

    assert Category.objects.count() == before
    assert Product.objects.get(code="VC-001").category.name == "Visitor chairs"


@pytest.mark.django_db
def test_duplicate_code_in_file_is_an_error(tmp_path):
    path = _xlsx(tmp_path, [
        ["VC-001", "Visitor chair", "Visitor chairs", "pcs", 2500],
        ["vc-001", "Same code again", "Visitor chairs", "pcs", 2600],
    ])

    with pytest.raises(CommandError) as exc:
        call_command("import_products", path)

    assert "already on row 2" in str(exc.value)
    assert not Product.objects.exists()


@pytest.mark.django_db
def test_template_has_headers_but_no_products(tmp_path):
    """Examples live on the help sheet, so an unfilled template can't create fake products."""
    path = str(tmp_path / "template.xlsx")

    call_command("import_products", template=path)
    call_command("import_products", path)

    assert not Product.objects.exists()


@pytest.mark.django_db
def test_filled_template_imports(tmp_path):
    from openpyxl import load_workbook

    path = str(tmp_path / "template.xlsx")
    call_command("import_products", template=path)
    wb = load_workbook(path)
    wb["Products"].append(["oc-001", "Office chair", "Office chairs", "pcs", 4200, 3900, 5,
                           ""])
    wb.save(path)

    call_command("import_products", path)

    product = Product.objects.get(code="OC-001")
    assert (product.selling_price, product.wholesale_price, product.min_stock) == (
        Decimal("4200.00"), Decimal("3900.00"), 5)


def _wholesale_xlsx(tmp_path, rows, name="w.xlsx"):
    wb = Workbook()
    ws = wb.active
    ws.append(["Code", "Name", "Category", "Unit", "Price", "Wholesale price"])
    for row in rows:
        ws.append(row)
    path = tmp_path / name
    wb.save(path)
    return str(path)


@pytest.mark.django_db
def test_wholesale_above_price_is_a_row_error(tmp_path):
    path = _wholesale_xlsx(tmp_path, [["VC-001", "Chair", "Visitor chairs", "pcs", 2500, 2600]])

    with pytest.raises(CommandError) as exc:
        call_command("import_products", path)

    assert "row 2" in str(exc.value)
    assert not Product.objects.exists()


@pytest.mark.django_db
def test_lowering_both_prices_applies_them_in_a_safe_order(tmp_path):
    call_command("import_products", _wholesale_xlsx(
        tmp_path, [["VC-001", "Chair", "Visitor chairs", "pcs", 2500, 2200]], "a.xlsx"))

    # The new price 2100 is below the old wholesale 2200, so wholesale must drop first.
    call_command("import_products", _wholesale_xlsx(
        tmp_path, [["VC-001", "Chair", "Visitor chairs", "pcs", 2100, 1900]], "b.xlsx"))

    product = Product.objects.get(code="VC-001")
    assert (product.selling_price, product.wholesale_price) == (Decimal("2100.00"),
                                                                 Decimal("1900.00"))
    assert set(PriceHistory.objects.values_list("price_type", flat=True)) == {
        "selling", "wholesale"}


@pytest.mark.django_db
def test_price_below_existing_wholesale_without_new_wholesale_is_reported(tmp_path):
    call_command("import_products", _wholesale_xlsx(
        tmp_path, [["VC-001", "Chair", "Visitor chairs", "pcs", 2500, 2200]], "a.xlsx"))

    with pytest.raises(CommandError) as exc:
        call_command("import_products", _wholesale_xlsx(
            tmp_path, [["VC-001", "Chair", "Visitor chairs", "pcs", 2000, None]], "b.xlsx"))

    assert "cannot be higher than the selling price" in str(exc.value)
    assert Product.objects.get(code="VC-001").selling_price == Decimal("2500.00")
