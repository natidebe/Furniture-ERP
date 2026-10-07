"""P-14: download the template, check it (dry run), then import it — over the API."""

from io import BytesIO

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import Workbook, load_workbook

from apps.catalog.models import Product

URL = "/api/v1/products/import/"


def _upload(rows, name="products.xlsx"):
    wb = Workbook()
    ws = wb.active
    ws.append(["Code", "Name", "Category", "Unit", "Price", "Wholesale price"])
    for row in rows:
        ws.append(row)
    buffer = BytesIO()
    wb.save(buffer)
    return SimpleUploadedFile(name, buffer.getvalue())


@pytest.mark.django_db
def test_check_then_import(client_for):
    client, admin = client_for("admin")
    rows = [["VC-001", "Visitor chair", "Visitor chairs", "pcs", 2500, 2200],
            ["DS-003", "Manager desk", "Desks", "pcs", 32000, None]]

    check = client.post(URL, {"file": _upload(rows), "dry_run": "true"}, format="multipart")
    assert check.status_code == 200, check.data
    assert (check.data["dry_run"], check.data["created"]) == (True, 2)
    assert not Product.objects.exists()

    done = client.post(URL, {"file": _upload(rows)}, format="multipart")
    assert done.status_code == 200, done.data
    assert (done.data["dry_run"], done.data["created"]) == (False, 2)
    assert Product.objects.get(code="VC-001").created_by == admin


@pytest.mark.django_db
def test_bad_rows_are_listed_and_nothing_is_imported(client_for):
    client, _ = client_for("admin")
    rows = [["VC-001", "Visitor chair", "Visitor chairs", "pcs", 2500, None],
            ["VC-002", "Bad unit", "Visitor chairs", "boxes", 100, None]]

    response = client.post(URL, {"file": _upload(rows)}, format="multipart")

    assert response.status_code == 400
    assert response.data["code"] == "import_failed"
    assert response.data["errors"] == ["row 3: unknown unit 'boxes'"]
    assert not Product.objects.exists()


@pytest.mark.django_db
def test_not_a_spreadsheet(client_for):
    client, _ = client_for("admin")

    wrong_type = client.post(URL, {"file": SimpleUploadedFile("p.csv", b"a,b")},
                             format="multipart")
    broken = client.post(URL, {"file": SimpleUploadedFile("p.xlsx", b"not a zip")},
                         format="multipart")

    assert wrong_type.status_code == 400 and "file" in wrong_type.data
    assert broken.status_code == 400
    assert broken.data["errors"] == ["The file is not an .xlsx spreadsheet."]


@pytest.mark.django_db
@pytest.mark.parametrize("role", ["accountant", "salesperson", "storekeeper"])
def test_only_admins_import(client_for, role):
    client, _ = client_for(role)

    response = client.post(URL, {"file": _upload([])}, format="multipart")
    template = client.get("/api/v1/products/import-template/")

    assert response.status_code == 403
    assert template.status_code == 403


@pytest.mark.django_db
def test_template_download_can_be_filled_and_imported(client_for):
    client, _ = client_for("admin")

    response = client.get("/api/v1/products/import-template/")
    assert response.status_code == 200
    assert "products-template.xlsx" in response["Content-Disposition"]

    wb = load_workbook(BytesIO(response.content))
    wb["Products"].append(["OC-001", "Office chair", "Office chairs", "pcs", 4200, 3900, 5,
                           ""])
    buffer = BytesIO()
    wb.save(buffer)
    upload = SimpleUploadedFile("filled.xlsx", buffer.getvalue())

    assert client.post(URL, {"file": upload}, format="multipart").data["created"] == 1
