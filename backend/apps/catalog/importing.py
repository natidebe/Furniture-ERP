"""Product import from an .xlsx file, shared by `manage.py import_products` and
POST /products/import/ (P-14).

Columns: code, name, category, unit, price, and optionally wholesale price, min stock and
description. Idempotent: existing codes are updated and price changes go through
change_price, so PriceHistory is kept. Any bad row aborts the whole import.
"""

from decimal import Decimal, InvalidOperation
from zipfile import BadZipFile

from django.db import transaction
from django.db.models import Q
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.worksheet.datavalidation import DataValidation

from apps.catalog import services
from apps.catalog.models import Category, PriceType, Product, Unit
from apps.core.exceptions import BusinessRuleError

COLUMNS = ("code", "name", "category", "unit", "price")
OPTIONAL_COLUMNS = ("wholesale price", "min stock", "description")
TEMPLATE_HEADERS = ["Code", "Name", "Category", "Unit", "Price", "Wholesale price", "Min stock",
                    "Description"]
TEMPLATE_EXAMPLES = [
    ("VC-001", "Visitor chair, black mesh", "Visitor chairs", "pcs", 2500, 2200, 10, ""),
    ("EC-010", "Executive chair, leather", "Executive chairs", "pcs", 18500, 16500, 2,
     "High back"),
    ("DS-003", "Manager desk 160 cm", "Desks", "pcs", 32000, None, 1, ""),
]


class ProductImportError(Exception):
    """Nothing was imported. `errors` lists each problem ("row 7: unknown unit 'boxes'")."""

    def __init__(self, errors: list[str]):
        super().__init__("Nothing imported. Fix these and try again:\n  "
                         + "\n  ".join(errors))
        self.errors = errors


def import_products(source, *, user=None, dry_run: bool = False) -> dict[str, int]:
    """Import `source` (a path or an open binary file). Returns counts of created, updated,
    price_changed and unchanged rows; with dry_run nothing is saved."""
    rows = list(_read_rows(source))
    counts = {"created": 0, "updated": 0, "price_changed": 0, "unchanged": 0}
    errors = []
    seen_codes: dict[str, int] = {}

    with transaction.atomic():
        for line_no, row in rows:
            code = str(row["code"] or "").strip().upper()
            if code and code in seen_codes:
                errors.append(f"row {line_no}: code {code} already on row {seen_codes[code]}")
                continue
            seen_codes[code] = line_no
            try:
                result = _import_row(row, user)
            except BusinessRuleError as exc:
                errors.append(f"row {line_no}: {exc.message}")
                continue
            except ValueError as exc:
                errors.append(f"row {line_no}: {exc}")
                continue
            for key in result:
                counts[key] += 1

        if errors:
            transaction.set_rollback(True)
            raise ProductImportError(errors)
        if dry_run:
            transaction.set_rollback(True)
    return counts


def _read_rows(source):
    try:
        sheet = load_workbook(source, read_only=True, data_only=True).active
    except FileNotFoundError as exc:
        raise ProductImportError([f"File not found: {source}"]) from exc
    except (InvalidFileException, BadZipFile, KeyError, OSError) as exc:
        raise ProductImportError(["The file is not an .xlsx spreadsheet."]) from exc
    rows = sheet.iter_rows(values_only=True)
    header = [str(cell or "").strip().lower() for cell in next(rows, ())]
    missing = [c for c in COLUMNS if c not in header]
    if missing:
        raise ProductImportError([f"Missing columns: {', '.join(missing)}"])
    index = {c: header.index(c) for c in COLUMNS + OPTIONAL_COLUMNS if c in header}
    for line_no, values in enumerate(rows, start=2):
        if not any(v not in (None, "") for v in values):
            continue
        yield line_no, {c: (values[i] if i < len(values) else None) for c, i in index.items()}


def _import_row(row, actor) -> list[str]:
    code = str(row["code"] or "").strip().upper()
    name = str(row["name"] or "").strip()
    if not code or not name:
        raise ValueError("code and name are required")
    price = _money(row["price"], "price")
    if price is None or price <= 0:
        raise ValueError("price must be greater than zero")
    # Blank wholesale: none for a new product, unchanged for an existing one.
    wholesale = _money(row.get("wholesale price"), "wholesale price")
    if wholesale is not None and not 0 < wholesale <= price:
        raise ValueError(f"wholesale price {wholesale} must be above zero and not above "
                         f"the price {price}")

    category_name = " ".join(str(row["category"] or "").split())
    if not category_name:
        raise ValueError("category is required")
    category = Category.objects.filter(name__iexact=category_name).first()
    if category is None:
        category = Category.objects.create(name=category_name, created_by=actor)

    unit_name = str(row["unit"] or "").strip()
    unit = Unit.objects.filter(Q(symbol__iexact=unit_name) | Q(name__iexact=unit_name)).first()
    if unit is None:
        raise ValueError(f"unknown unit {unit_name!r}")

    details = {"name": name, "category": category, "unit": unit}
    if "min stock" in row:
        details["min_stock"] = _min_stock(row["min stock"])
    if "description" in row:
        details["description"] = str(row["description"] or "").strip()

    product = Product.objects.filter(code=code).first()
    if product is None:
        services.create_product(user=actor, code=code, selling_price=price,
                                wholesale_price=wholesale, **details)
        return ["created"]

    result = []
    if any(getattr(product, k) != v for k, v in details.items()):
        product = services.update_product(user=actor, product=product, **details)
        result.append("updated")
    changes = []
    if product.selling_price != price:
        changes.append((PriceType.SELLING, price))
    if wholesale is not None and product.wholesale_price != wholesale:
        changes.append((PriceType.WHOLESALE, wholesale))
    if wholesale is not None and price < product.selling_price:
        # Lower the wholesale price first so it never sits above the selling price.
        changes.reverse()
    for price_type, value in changes:
        product = services.change_price(product=product, new_price=value, user=actor,
                                        price_type=price_type, reason="import_products")
    if changes:
        result.append("price_changed")
    return result or ["unchanged"]


def _money(value, label) -> Decimal | None:
    if value in (None, ""):
        return None
    try:
        return Decimal(str(value)).quantize(Decimal("0.01"))
    except (InvalidOperation, TypeError) as exc:
        raise ValueError(f"invalid {label} {value!r}") from exc


def _min_stock(value) -> int:
    if value in (None, ""):
        return 0
    try:
        number = Decimal(str(value))
    except InvalidOperation as exc:
        raise ValueError(f"invalid min stock {value!r}") from exc
    if number < 0 or number != number.to_integral_value():
        raise ValueError(f"min stock must be a whole number ≥ 0, got {value!r}")
    return int(number)


def write_template(dest) -> None:
    """Write a blank template to `dest` (a path or an open binary file): the Products sheet,
    a Lists sheet with categories and units, and a How to fill sheet with examples."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Products"
    ws.append(TEMPLATE_HEADERS)
    bold, fill = Font(bold=True), PatternFill("solid", fgColor="DDEBF7")
    for cell in ws[1]:
        cell.font, cell.fill = bold, fill
    for col, width in zip("ABCDEFGH", (12, 36, 24, 8, 12, 16, 11, 30), strict=True):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    for row in range(2, 2001):
        ws[f"E{row}"].number_format = "#,##0.00"
        ws[f"F{row}"].number_format = "#,##0.00"

    lists = wb.create_sheet("Lists")
    lists.append(["Categories (new names are added on import)", "Units"])
    for cell in lists[1]:
        cell.font = bold
    categories = list(Category.objects.filter(is_active=True).values_list("name", flat=True))
    units = list(Unit.objects.values_list("symbol", flat=True))
    for i in range(max(len(categories), len(units))):
        lists.append([categories[i] if i < len(categories) else None,
                      units[i] if i < len(units) else None])
    lists.column_dimensions["A"].width = 42

    unit_rule = DataValidation(type="list", formula1=f"=Lists!$B$2:$B${len(units) + 1}",
                               allow_blank=False, showErrorMessage=True,
                               errorTitle="Unknown unit", error="Pick a unit from the list.")
    category_hint = DataValidation(type="list",
                                   formula1=f"=Lists!$A$2:$A${len(categories) + 1}",
                                   allow_blank=False, showErrorMessage=False)
    ws.add_data_validation(unit_rule)
    ws.add_data_validation(category_hint)
    unit_rule.add("D2:D2000")
    category_hint.add("C2:C2000")

    notes = wb.create_sheet("How to fill")
    for line in [
        "Fill the Products sheet: one row per product, starting on row 2.",
        "Code: unique, e.g. VC-001. It is stored in capitals.",
        "Price: selling price in ETB, greater than zero. No cost price (not recorded).",
        "Wholesale price: what resellers pay, in ETB. Optional; not above the price. "
        "Empty: resellers pay the normal price (an existing product keeps its old one).",
        "Unit: pick from the list (pcs or set).",
        "Category: pick from the list, or type a new one.",
        "Min stock: optional. A low-stock alert is sent when total stock goes below it.",
        "Description: optional.",
        "",
        "Examples (do not copy these into the Products sheet unless they are real):",
    ]:
        notes.append([line])
    notes.append(TEMPLATE_HEADERS)
    for cell in notes[notes.max_row]:
        cell.font = bold
    for row in TEMPLATE_EXAMPLES:
        notes.append(row)
    notes.column_dimensions["A"].width = 90
    wb.save(dest)
