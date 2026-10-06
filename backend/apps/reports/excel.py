"""One Excel workbook per report: formatted headers, ETB number format, a totals row,
frozen header and sized columns (BUILD_PHASES.md 4.3)."""

from datetime import date, datetime
from decimal import Decimal
from io import BytesIO

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

MONEY = "#,##0.00"
HEADER_FILL = PatternFill("solid", fgColor="DDEBF7")
BOLD = Font(bold=True)


def _cell(value):
    if isinstance(value, datetime):
        return timezone.localtime(value).replace(tzinfo=None) if timezone.is_aware(value) \
            else value
    return value


def _money(value):
    return Decimal(value) if value not in (None, "") else None


def _table(ws, headers, rows, *, money=(), totals=(), start_row=1):
    """headers = [(label, key)]; rows = list of dicts. `money` / `totals` are keys."""
    for col, (label, _) in enumerate(headers, start=1):
        cell = ws.cell(row=start_row, column=col, value=label)
        cell.font, cell.fill = BOLD, HEADER_FILL
    for r, row in enumerate(rows, start=start_row + 1):
        for col, (_, key) in enumerate(headers, start=1):
            value = row.get(key)
            cell = ws.cell(row=r, column=col,
                           value=_money(value) if key in money else _cell(value))
            if key in money:
                cell.number_format = MONEY
            elif isinstance(cell.value, datetime):
                cell.number_format = "dd/mm/yyyy hh:mm"
    if totals and rows:
        r = start_row + len(rows) + 1
        ws.cell(row=r, column=1, value="Total").font = BOLD
        for col, (_, key) in enumerate(headers, start=1):
            if key in totals:
                values = [row.get(key) for row in rows if row.get(key) not in (None, "")]
                total = sum((Decimal(str(v)) for v in values), Decimal("0"))
                cell = ws.cell(row=r, column=col, value=total)
                cell.font = BOLD
                if key in money:
                    cell.number_format = MONEY
    ws.freeze_panes = ws.cell(row=start_row + 1, column=1)
    for col, (label, key) in enumerate(headers, start=1):
        longest = max([len(str(label))] + [len(str(row.get(key) or "")) for row in rows[:500]])
        ws.column_dimensions[get_column_letter(col)].width = min(max(10, longest + 2), 50)


def _summary(ws, pairs, money=()):
    ws.append(["Item", "Value"])
    for cell in ws[1]:
        cell.font, cell.fill = BOLD, HEADER_FILL
    for label, key, value in pairs:
        ws.append([label, _money(value) if key in money and value is not None else value])
        if key in money:
            ws.cell(row=ws.max_row, column=2).number_format = MONEY
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 22


def _sales(wb, data):
    ws = wb.active
    ws.title = "Summary"
    received = data["received"]
    pairs = [("From", "from", data["from"]), ("To", "to", data["to"]),
             ("Sales", "sales", data["sales"]), ("Returns", "returns", data["returns"]),
             ("Net sales", "net", data["net_sales"]),
             ("Transactions", "tx", data["transactions"]),
             ("Paid (of these sales)", "paid", data.get("paid")),
             ("Credit (still owed)", "credit", data.get("credit")),
             ("Official receipt sales", "official", data["official_receipt"]),
             ("No-receipt sales", "none", data["no_receipt"]),
             ("Received — Organization", "org", received["organization"])]
    if received["personal"] is not None:
        pairs += [("Received — Personal", "per", received["personal"]),
                  ("Received — Combined", "comb", received["combined"])]
    _summary(ws, pairs, money={"sales", "returns", "net", "paid", "credit", "official", "none",
                               "org", "per", "comb"})
    _table(wb.create_sheet("By salesperson"),
           [("Salesperson", "salesperson"), ("Sales (ETB)", "sales"),
            ("Transactions", "transactions")], data["by_salesperson"], money={"sales"},
           totals={"sales", "transactions"})
    _table(wb.create_sheet("By branch"),
           [("Branch", "branch"), ("Sales (ETB)", "sales"), ("Transactions", "transactions")],
           data["by_branch"], money={"sales"}, totals={"sales", "transactions"})
    _table(wb.create_sheet("Products"),
           [("Code", "code"), ("Product", "name"), ("Qty", "qty"), ("Sales (ETB)", "sales")],
           data["products"], money={"sales"}, totals={"qty", "sales"})
    if data.get("by_month"):
        _table(wb.create_sheet("By month"),
               [("Month", "month"), ("Sales (ETB)", "sales"), ("Returns (ETB)", "returns"),
                ("Net sales (ETB)", "net_sales"), ("Transactions", "transactions")],
               data["by_month"], money={"sales", "returns", "net_sales"},
               totals={"sales", "returns", "net_sales", "transactions"})


def _payments(wb, data):
    ws = wb.active
    ws.title = "Totals"
    t = data["totals"]
    pairs = [("From", "from", data["from"]), ("To", "to", data["to"]),
             ("Organization", "org", t["organization"])]
    if t["personal"] is not None:
        pairs += [("Personal", "per", t["personal"]), ("Combined", "comb", t["combined"])]
    _summary(ws, pairs, money={"org", "per", "comb"})
    headers = [("Period", "period"), ("Organization (ETB)", "organization")]
    if t["personal"] is not None:
        headers += [("Personal (ETB)", "personal"), ("Combined (ETB)", "combined")]
    _table(wb.create_sheet("By period"), headers, data["by_period"],
           money={"organization", "personal", "combined"},
           totals={"organization", "personal", "combined"})
    _table(wb.create_sheet("Payments"),
           [("Number", "number"), ("Paid at", "paid_at"), ("Customer", "customer"),
            ("Amount (ETB)", "amount"), ("Account", "account"), ("Kind", "kind"),
            ("Method", "method"), ("Receipt", "receipt_number"), ("Status", "status"),
            ("Recorded by", "recorded_by")], data["payments"], money={"amount"},
           totals={"amount"})


def _credit(wb, data):
    ws = wb.active
    ws.title = "Customers"
    _table(ws, [("Customer", "customer"), ("Shop", "shop_name"), ("Phone", "phone"),
                ("City", "city"), ("Credit limit", "credit_limit"),
                ("Outstanding (ETB)", "outstanding"), ("0–30 days", "0_30"),
                ("31–60 days", "31_60"), ("61–90 days", "61_90"), ("Over 90 days", "over_90")],
           data["customers"],
           money={"credit_limit", "outstanding", "0_30", "31_60", "61_90", "over_90"},
           totals={"outstanding", "0_30", "31_60", "61_90", "over_90"})
    _summary(wb.create_sheet("Summary"),
             [("From", "from", data["from"]), ("To", "to", data["to"]),
              ("Outstanding total", "out", data["outstanding_total"]),
              ("Credit collected", "col", data["credit_collected"])], money={"out", "col"})


def _stock(wb, data):
    ws = wb.active
    ws.title = "Stock"
    headers = [("Code", "code"), ("Product", "name"), ("Unit", "unit")]
    rows = []
    for product in data["products"]:
        row = dict(product)
        for code in data["locations"]:
            row[f"loc_{code}"] = product["stock"].get(code, 0)
        rows.append(row)
    headers += [(code if code != "TRANSIT" else "In transit", f"loc_{code}")
                for code in data["locations"]]
    headers += [("Total", "total"), ("Display", "display"), ("Damaged", "damaged"),
                ("Sellable", "total_new"), ("Minimum", "min_stock"), ("Low", "low_stock")]
    _table(ws, headers, rows, totals={f"loc_{c}" for c in data["locations"]}
           | {"total", "display", "damaged", "total_new"})


def _list(name, headers, key, money=(), totals=()):
    def build(wb, data):
        ws = wb.active
        ws.title = name
        _table(ws, headers, data[key], money=money, totals=totals)
    return build


BUILDERS = {
    "sales": _sales,
    "payments": _payments,
    "credit": _credit,
    "stock": _stock,
    "movements": _list("Movements", [
        ("Number", "number"), ("Date", "occurred_at"), ("Type", "type"),
        ("Condition", "condition"), ("Product", "product"), ("Qty", "qty"), ("From", "from"),
        ("To", "to"), ("Customer", "customer"), ("Transaction", "transaction"),
        ("Reference", "reference"), ("Person", "person")], "movements"),
    "open-requests": _list("Open requests", [
        ("Number", "number"), ("Status", "status"), ("Branch", "branch"), ("Source", "source"),
        ("Customer", "customer"), ("Salesperson", "salesperson"), ("Created", "created_at"),
        ("Waiting (hours)", "waiting_hours"), ("Units left", "units_remaining")], "requests"),
    "unverified-payments": _list("Unverified payments", [
        ("Number", "number"), ("Paid at", "paid_at"), ("Customer", "customer"),
        ("Amount (ETB)", "amount"), ("Account", "account"), ("Kind", "kind"),
        ("Receipt", "receipt_number"), ("Recorded by", "recorded_by")], "payments",
        money={"amount"}, totals={"amount"}),
}


def workbook_bytes(name: str, data: dict) -> bytes:
    wb = Workbook()
    BUILDERS[name](wb, data)
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def filename(name: str, data: dict) -> str:
    first, last = data.get("from"), data.get("to")
    if isinstance(first, date) and isinstance(last, date):
        span = first.isoformat() if first == last else f"{first.isoformat()}_{last.isoformat()}"
    else:
        span = timezone.localdate().isoformat()
    return f"{name}_{span}.xlsx"
