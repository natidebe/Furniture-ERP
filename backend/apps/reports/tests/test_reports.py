from datetime import timedelta
from decimal import Decimal
from io import BytesIO

import pytest
from django.utils import timezone
from openpyxl import load_workbook

from apps.accounts.services import set_erp_permissions
from apps.customers.selectors import customer_balance
from apps.notifications.models import NotificationOutbox
from apps.payments import services as payments
from apps.reports import selectors, tasks
from apps.sales import services as sales
from apps.sales.models import SalesOrder

D = Decimal


@pytest.fixture
def day(staff, loc, goods, abc, walk_in, accounts, stock, make_user):
    """Today: two sales at Piassa (one official, paid; one on credit), one at Denbel."""
    for product in (goods.chair, goods.desk, goods.cabinet):
        stock(product, loc.PIA, 20)
        stock(product, loc.DEN, 20)
    den = make_user(role="salesperson", home_location=loc.DEN, full_name="Sara")
    staff.sales.full_name = "Abebe"
    staff.sales.save()

    official = sales.create_order(customer=walk_in, branch=loc.PIA, user=staff.sales,
                                  receipt_type="official",
                                  lines=[{"product": goods.desk, "qty": 2}])  # 40,000
    payments.record_payment(customer=walk_in, account=accounts.org, amount="40000",
                            method="bank", recorded_by=staff.accountant, receipt_number="R-1",
                            allocations=[{"order": official, "amount": "40000"}])
    sales.confirm_order(order=official, user=staff.sales)

    credit = sales.confirm_order(order=sales.create_order(
        customer=abc, branch=loc.PIA, user=staff.sales,
        lines=[{"product": goods.chair, "qty": 10}]), user=staff.sales)  # 45,000 wholesale
    payments.record_payment(customer=abc, account=accounts.personal, amount="15000",
                            method="cash", recorded_by=staff.accountant,
                            allocations=[{"order": credit, "amount": "15000"}])

    denbel = sales.confirm_order(order=sales.create_order(
        customer=abc, branch=loc.DEN, user=den,
        lines=[{"product": goods.cabinet, "qty": 1}]), user=den)  # 20,000
    return {"official": official, "credit": credit, "denbel": denbel, "den": den}


def _today():
    today = timezone.localdate()
    return today, today


@pytest.mark.django_db
def test_daily_sales_figures(day, staff):
    report = selectors.sales_report(staff.admin, *_today())

    assert report["sales"] == "105000.00" == report["net_sales"]
    assert report["transactions"] == 3
    assert (report["paid"], report["credit"]) == ("55000.00", "50000.00")
    assert (report["official_receipt"], report["no_receipt"]) == ("40000.00", "65000.00")
    assert {r["branch"]: r["sales"] for r in report["by_branch"]} == {
        "PIA": "85000.00", "DEN": "20000.00"}
    assert {r["salesperson"]: r["sales"] for r in report["by_salesperson"]} == {
        "Abebe": "85000.00", "Sara": "20000.00"}
    assert report["best_sellers"][0] == {"code": "VC-001", "name": "Visitor chair",
                                         "qty": 10, "sales": "45000.00"}
    assert report["received"] == {"organization": "40000.00", "personal": "15000.00",
                                  "combined": "55000.00"}


@pytest.mark.django_db
def test_sales_total_equals_confirmed_order_totals(day, staff):
    report = selectors.sales_report(staff.admin, *_today())
    total = sum(o.total_amount for o in SalesOrder.objects.filter(confirmed_at__isnull=False))

    assert D(report["sales"]) == total


@pytest.mark.django_db
def test_returns_count_in_the_period_they_happen(day, staff, loc):
    yesterday = timezone.localdate() - timedelta(days=1)
    SalesOrder.objects.filter(pk=day["credit"].pk).update(
        confirmed_at=timezone.now() - timedelta(days=1))
    sales.return_goods(order=day["credit"], user=staff.accountant, location=loc.PIA,
                       reason="Wrong colour",
                       lines=[{"line_id": day["credit"].lines.get().pk, "qty": 2}])

    before = selectors.sales_report(staff.admin, yesterday, yesterday)
    today = selectors.sales_report(staff.admin, *_today())

    assert before["sales"] == "45000.00" and before["returns"] == "0.00"  # unchanged
    assert today["returns"] == "9000.00"
    assert today["net_sales"] == "51000.00"  # 60,000 sold today − 9,000 returned


@pytest.mark.django_db
def test_salesperson_report_shows_only_own_sales(day, staff):
    report = selectors.sales_report(staff.sales, *_today())

    assert report["transactions"] == 2
    assert [r["salesperson"] for r in report["by_salesperson"]] == ["Abebe"]


@pytest.mark.django_db
def test_payments_report_organization_plus_personal_is_combined(day, staff):
    report = selectors.payments_report(staff.admin, *_today())
    t = report["totals"]

    assert D(t["organization"]) + D(t["personal"]) == D(t["combined"]) == D("55000")
    assert len(report["payments"]) == 2
    assert report["by_period"][0]["combined"] == "55000.00"


@pytest.mark.django_db
def test_payments_report_by_order_and_without_personal_permission(day, staff, make_user):
    by_order = selectors.payments_report(staff.admin, *_today(),
                                         filters={"order": day["credit"].pk})
    assert [p["kind"] for p in by_order["payments"]] == ["personal"]

    restricted = make_user(role="accountant")
    set_erp_permissions(restricted, [])
    hidden = selectors.payments_report(restricted, *_today())
    assert hidden["totals"] == {"organization": "40000.00", "personal": None, "combined": None}
    assert {p["kind"] for p in hidden["payments"]} == {"organization"}


@pytest.mark.django_db
def test_credit_report_matches_customer_balances(day, staff, abc):
    report = selectors.credit_report(staff.admin, *_today())

    assert D(report["outstanding_total"]) == customer_balance(abc)["outstanding"] == D("50000")
    row = report["customers"][0]
    assert row["customer"] == "ABC Furniture"
    assert row["0_30"] == "50000.00" and row["over_90"] == "0.00"


@pytest.mark.django_db
def test_credit_ageing_and_collection(day, staff, abc, accounts):
    SalesOrder.objects.filter(pk=day["credit"].pk).update(
        confirmed_at=timezone.now() - timedelta(days=45))

    payments.record_payment(customer=abc, account=accounts.org, amount="5000", method="bank",
                            recorded_by=staff.accountant,
                            allocations=[{"order": day["credit"], "amount": "5000"}])
    report = selectors.credit_report(staff.admin, *_today())

    row = report["customers"][0]
    assert (row["31_60"], row["0_30"]) == ("25000.00", "20000.00")
    # Paid today on a sale confirmed 45 days ago: today's 15,000 + 5,000.
    assert report["credit_collected"] == "20000.00"


@pytest.mark.django_db
def test_yearly_months_add_up(day, staff):
    first, last = selectors.period_range("year")
    report = selectors.sales_report(staff.admin, first, last)

    assert len(report["by_month"]) == 12
    assert sum(D(m["sales"]) for m in report["by_month"]) == D(report["sales"])


def test_period_ranges():
    from datetime import date

    wednesday = date(2026, 10, 7)
    assert selectors.period_range("week", wednesday) == (date(2026, 10, 5), date(2026, 10, 11))
    assert selectors.period_range("month", date(2026, 2, 10)) == (date(2026, 2, 1),
                                                                  date(2026, 2, 28))
    assert selectors.period_range("year", wednesday) == (date(2026, 1, 1), date(2026, 12, 31))


# ---------------------------------------------------------------- API and Excel

@pytest.mark.django_db
def test_report_api_json_and_excel(day, client_for):
    client, _ = client_for("accountant")

    data = client.get("/api/v1/reports/sales/", {"period": "day"})
    assert data.status_code == 200 and data.data["transactions"] == 3

    xlsx = client.get("/api/v1/reports/sales/", {"period": "day", "format": "xlsx"})
    assert xlsx.status_code == 200
    assert "attachment" in xlsx["Content-Disposition"]
    wb = load_workbook(BytesIO(xlsx.content))
    assert wb.sheetnames == ["Summary", "By salesperson", "By branch", "Products"]
    products = wb["Products"]
    assert products["A1"].value == "Code"
    assert products.cell(row=products.max_row, column=1).value == "Total"
    assert products.cell(row=products.max_row, column=4).value == D("105000.00")


@pytest.mark.django_db
@pytest.mark.parametrize("name", ["sales", "payments", "credit", "stock", "movements",
                                  "open-requests", "unverified-payments"])
def test_every_report_exports(day, client_for, name):
    client, _ = client_for("admin")

    response = client.get(f"/api/v1/reports/{name}/", {"period": "month", "format": "xlsx"})

    assert response.status_code == 200, response.content[:200]
    assert load_workbook(BytesIO(response.content)).sheetnames


@pytest.mark.django_db
def test_excel_needs_export_permission(day, client_for):
    client, user = client_for("accountant")
    set_erp_permissions(user, ["view_personal_payments"])

    assert client.get("/api/v1/reports/sales/", {"format": "xlsx"}).status_code == 403
    assert client.get("/api/v1/reports/sales/").status_code == 200


@pytest.mark.django_db
def test_report_access_by_role(day, client_for, loc):
    sales_client, _ = client_for("salesperson", home_location=loc.PIA)
    assert sales_client.get("/api/v1/reports/sales/").status_code == 200
    assert sales_client.get("/api/v1/reports/payments/").status_code == 403

    store, _ = client_for("storekeeper", home_location=loc.PAW)
    assert store.get("/api/v1/reports/stock/").status_code == 200
    assert store.get("/api/v1/reports/sales/").status_code == 403
    assert store.get("/api/v1/reports/nope/").status_code == 404


@pytest.mark.django_db
def test_bad_dates_are_a_clear_error(client_for):
    client, _ = client_for("admin")

    assert client.get("/api/v1/reports/sales/", {"from": "2026-10-05"}).status_code == 400
    assert client.get("/api/v1/reports/sales/", {"period": "decade"}).status_code == 400


# ---------------------------------------------------------------- scheduled reports

@pytest.mark.django_db
def test_daily_report_message_to_the_owner(day, staff):
    staff.admin.telegram_id = 1
    staff.admin.save()

    assert tasks.daily_report() == 1

    text = NotificationOutbox.objects.get(event_type="report.daily").payload["text"]
    assert "Daily Sales" in text
    assert "Total sales: 105,000 ETB (3 sales)" in text
    assert "Paid: 55,000 ETB · Credit: 50,000 ETB" in text
    assert "Organization: 40,000 ETB · Personal: 15,000 ETB" in text
    assert "VC-001 ×10" in text


@pytest.mark.django_db
def test_weekly_report_has_credit_and_an_excel_file(day, staff):
    staff.admin.telegram_id = 1
    staff.admin.save()

    tasks.weekly_report()

    payload = NotificationOutbox.objects.get(event_type="report.weekly").payload
    assert "Outstanding credit (all customers): 50,000 ETB" in payload["text"]
    assert payload["document"]["filename"].startswith("sales_")
    assert payload["document"]["content_b64"]
